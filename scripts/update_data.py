#!/usr/bin/env python3
"""Export the Village Board Tracker spreadsheet to data/meetings.json for the website.

Usage:
    python3 scripts/update_data.py                      # default tracker path, redaction on
    python3 scripts/update_data.py --tracker other.xlsx
    python3 scripts/update_data.py --show-redactions    # print each address that was removed
    python3 scripts/update_data.py --no-redact-public   # keep addresses (not for publishing)

Same export logic as export_tracker.py, plus an "updated" timestamp and
optional removal of street addresses from public comment text.
"""
import argparse
import json
import os
import re
import sys
from datetime import datetime, timezone

try:
    import openpyxl
except ImportError:
    sys.exit("This script needs openpyxl. Install it with:\n    pip3 install --user -r requirements.txt")

DEFAULT_TRACKER = os.path.expanduser(
    "~/Library/Mobile Documents/com~apple~CloudDocs/Village Board/Village Board Tracker.xlsx"
)
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_OUT = os.path.join(ROOT, "data", "meetings.json")

# --- address redaction -------------------------------------------------------
# A street address here is a house number followed by an optional direction and
# one to four capitalised words, e.g. "815 Foxanna", "1100 Bayview Ct",
# "700 NW Hwy", "210 N River". Numbers followed by lowercase words, hyphens or
# punctuation ("5-0", "$15,000", "Troop 166", "6 years") are left alone, and so
# are all-caps words ("3 FOIA requests").
_DIR = r"(?:(?:N|S|E|W|NE|NW|SE|SW|North|South|East|West)\.?\s+)"
_WORD = r"(?:(?:St|Ave|Rd|Ct|Ln|Dr|Hwy|Blvd|Pl|Pkwy|Ter|Cir)\.|[A-Z][a-z][A-Za-z'-]*)"
ADDRESS = re.compile(rf"\b\d{{1,5}}[A-Za-z]?\s+{_DIR}?{_WORD}(?:\s+{_WORD}){{0,3}}")
YEAR = re.compile(r"^(19|20)\d\d\s")
PAREN = re.compile(r"\s*\(([^()]*)\)")
PLACEHOLDER = "[address removed]"


def _is_address(m):
    # "2024 Ford Explorer" is a year, not an address.
    return not YEAR.match(m.group(0))


def redact_addresses(text, log=None):
    """Remove street addresses from text. Returns the cleaned text."""
    if not text:
        return text

    def paren(m):
        inner = m.group(1)
        hits = [a for a in ADDRESS.finditer(inner) if _is_address(a)]
        if not hits:
            return m.group(0)
        if log is not None:
            log.extend(a.group(0) for a in hits)
        # "(815 Foxanna)" -> drop the whole parenthetical.
        if len(hits) == 1 and hits[0].group(0).strip() == inner.strip():
            return ""
        # "(renter behind 700 NW Hwy)" -> keep the context, hide the address.
        lead = m.group(0)[: len(m.group(0)) - len(m.group(0).lstrip())]
        return lead + "(" + ADDRESS.sub(lambda a: PLACEHOLDER if _is_address(a) else a.group(0), inner) + ")"

    text = PAREN.sub(paren, text)

    def bare(a):
        if not _is_address(a):
            return a.group(0)
        if log is not None:
            log.append(a.group(0))
        return PLACEHOLDER

    return ADDRESS.sub(bare, text)


# --- export ----------------------------------------------------------------
def export(tracker):
    wb = openpyxl.load_workbook(tracker)
    rows = []
    for r in wb["Meetings"].iter_rows(min_row=2):
        v = [c.value for c in r[:11]]
        if not v[0]:
            continue
        links = [(c.hyperlink.target if c.hyperlink else None) for c in r[7:10]]
        rows.append(dict(
            d=str(v[0])[:10], t=v[1] or "", s=v[2] or "", a=v[3] or "", o=v[4] or "",
            p=v[5] or "", g=v[6] or "", ag=links[0], pk=links[1], mn=links[2], k=v[10] or "",
        ))
    topics = [[str(c.value or "") for c in r[:4]] for r in wb["Topics"].iter_rows(min_row=2) if r[0].value]
    about = [[str(c.value or "") for c in r[:2]] for r in wb["About"].iter_rows() if r[0].value]
    return rows, topics, about


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--tracker", default=DEFAULT_TRACKER, help="path to Village Board Tracker.xlsx (default: %(default)s)")
    ap.add_argument("--out", default=DEFAULT_OUT, help="where to write the JSON (default: data/meetings.json)")
    ap.add_argument("--redact-public", action=argparse.BooleanOptionalAction, default=True,
                    help="strip street addresses from public comment text (default: on)")
    ap.add_argument("--show-redactions", action="store_true", help="print every address that was removed")
    args = ap.parse_args()

    if not os.path.exists(args.tracker):
        sys.exit(f"Can't find the tracker spreadsheet at:\n    {args.tracker}\nPass the right path with --tracker.")

    rows, topics, about = export(args.tracker)

    removed = []
    if args.redact_public:
        for r in rows:
            log = []
            r["p"] = redact_addresses(r["p"], log)
            if log and args.show_redactions:
                print(f"  {r['d']}: removed {', '.join(log)}")
            removed.extend(log)

    data = dict(
        updated=datetime.now(timezone.utc).isoformat(timespec="seconds"),
        redacted=bool(args.redact_public),
        rows=rows, topics=topics, about=about,
    )
    os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
    with open(args.out, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=1)
        f.write("\n")

    print(f"Exported {len(rows)} meetings, {len(topics)} storylines to {os.path.relpath(args.out)}")
    if args.redact_public:
        print(f"Removed {len(removed)} street addresses from public comments"
              + ("" if args.show_redactions else " (use --show-redactions to list them)"))
    else:
        print("WARNING: addresses were NOT removed. Don't publish this file.")


if __name__ == "__main__":
    main()
